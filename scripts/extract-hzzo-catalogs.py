"""Extract versioned HZZO reference workbooks. Dependencies: openpyxl, xlrd.

Never execute workbook formulas/macros. Effective dates win over revision dates.
Unversioned and legacy references are explicitly labelled; namespaces stay separate.
"""
import argparse, collections, datetime, gzip, hashlib, json, re
from pathlib import Path
import openpyxl, xlrd

TODAY = datetime.date(2026, 10, 11)

def text(value):
    if value is None: return ''
    if isinstance(value, float) and value.is_integer(): return str(int(value))
    return re.sub(r'\s+', ' ', str(value)).strip()

def load(path):
    if path.suffix.lower() == '.xlsx':
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheets = [(s.title, [list(row) for row in s.iter_rows(max_col=min(s.max_column or 100, 100), values_only=True)]) for s in book]
        book.close()
    else:
        book = xlrd.open_workbook(path)
        sheets = []
        for s in book.sheets():
            rows = []
            for i in range(s.nrows):
                rows.append([xlrd.xldate.xldate_as_datetime(c.value, book.datemode) if c.ctype == xlrd.XL_CELL_DATE else c.value for c in s.row(i)])
            sheets.append((s.name, rows))
    return [(name, [[v.isoformat() if isinstance(v,(datetime.date,datetime.datetime)) else v for v in row] for row in rows]) for name,rows in sheets]

def identify(name):
    # Fixed namespaces are intentional: billing groups are not procedure codes.
    if name.startswith('web_OLL'): return 'OLL', 'Osnovna lista lijekova'
    if name.startswith('web_DLL'): return 'DLL', 'Dopunska lista lijekova'
    if name.startswith('DTP_PZZ'): return 'DTP-PZZ', 'DTP postupci u primarnoj zdravstvenoj zaštiti'
    if re.search(r'DTP.*SKZZ',name): return 'DTP-SKZZ', 'DTP postupci u specijalističkoj zdravstvenoj zaštiti'
    if re.search(r'DTS_skupine',name): return 'DTS-SKUPINE', 'Dijagnostičko-terapijske skupine'
    if 'DBL_' in name: return 'DBL', 'Dugotrajno bolničko liječenje'
    if re.search(r'NPP',name): return 'NPP', 'Nacionalni preventivni programi'
    if 'Dobrovoljno davalstvo' in name: return 'KRV', 'Dobrovoljno davalaštvo krvi'
    if name.startswith('djelzz'): return 'DJELATNOSTI', 'Medicinske djelatnosti'
    if name.startswith('Medicinske djelatnosti'): return 'DJELATNOSTI-NEDATIRANO', 'Medicinske djelatnosti — nedatirana alternativna lista'
    if name.startswith('Aktivnosti'): return 'AKTIVNOSTI', 'Aktivnosti zdravstvene zaštite'
    if name.startswith('davzu'): return 'USTANOVE', 'Davatelji zdravstvenih usluga'
    if name.startswith('Naselja'): return 'NASELJA', 'Naselja i poštanski brojevi u RH'
    if name.startswith('Obloge'): return 'OBLOGE', 'Obloge za rane — dimenzije'
    if name.startswith('sudjelovanje'): return 'SUDJELOVANJE', 'Sudjelovanje u troškovima zdravstvene zaštite'
    if name.startswith('Dijagnoze_kod'): return 'OSLOBADANJE-MKB', 'Dijagnoze s punim pokrićem troškova zdravstvene zaštite'
    if name.startswith('Kategorije'): return 'KATEGORIJE', 'Kategorije osiguranika'
    if name.startswith('DZO_'): return 'DZO', 'Osiguravatelji dopunskog zdravstvenog osiguranja'
    if name.startswith('razdoz'): return 'RAZDOZ', 'Razlozi privremene nesposobnosti za rad'
    if name.startswith('SKZZ-Područje'): return 'PODRUCJA', 'SKZZ područja rada'
    if name.startswith('Pogreske_PZZ'): return 'GRESKE-PZZ', 'Pogreške PZZ predkontrole'
    if name.startswith('Šifre grešaka'): return 'GRESKE-RACUNA', 'Pogreške validacije specifikacija računa'
    if name.startswith('DTS_postupci'): return 'DTS-POSTUPCI', 'DTS postupci — referenca iz 2016.'
    if name.startswith('zdrpos'): return 'ZDRPOS', 'Zdravstveni postupci s bodovima — stara lista'
    return None

def date_info(file, family, sheets):
    if family == 'DZO': return None, 'Nepotvrđeno: naziv datoteke sadrži 2029.; datumi servisa nisu datum primjene cijele liste.', 'UNCONFIRMED'
    if family in ['OLL','DLL']:
        m=re.search(r'(\d{2})_(\d{2})_(\d{4})',file)
        if not m: raise ValueError('Missing medicine-list effective date')
        return f'{m[3]}-{m[2]}-{m[1]}', 'Datum stupanja na snagu iz naziva datoteke.', 'EFFECTIVE'
    if family == 'ZDRPOS': return '2011-11-30', 'Oznaka 111130 u nazivu; povijesna referenca, aktualnost nije potvrđena.', 'SNAPSHOT'
    if family == 'OBLOGE':
        title=sheets[0][0];m=re.search(r'(\d{2})\.(\d{2})\.(\d{4})',title)
        return f'{m[3]}-{m[2]}-{m[1]}', 'Datum iz naziva radnog lista.', 'SNAPSHOT'
    if family == 'NASELJA': return '2022-03-09', 'Datum iz naziva radnog lista.', 'SNAPSHOT'
    if family == 'USTANOVE':
        m=re.search(r'(\d{2})\.(\d{2})\.(\d{4})',file)
        return f'{m[3]}-{m[2]}-{m[1]}', 'Datum objave iz naziva datoteke; list aktivne.', 'SNAPSHOT'
    if family == 'DTP-PZZ':
        title=' '.join(text(v) for _,rows in sheets for row in rows[:4] for v in row)
        match=re.search(r'od\s+(\d{1,2})\.(\d{1,2})\.(\d{4})',title)
        if match: return datetime.date(int(match[3]),int(match[2]),int(match[1])).isoformat(), 'Datum primjene iz zaglavlja; datum revizije datoteke ne mijenja početak primjene.', 'EFFECTIVE'
    m=re.search(r'(\d{1,2})\.(\d{1,2})\.(\d{4})',file)
    if m: return datetime.date(int(m[3]),int(m[2]),int(m[1])).isoformat(), 'Datum iz naziva datoteke.', 'SNAPSHOT' if family in ['DTS-POSTUPCI','OSLOBADANJE-MKB','GRESKE-PZZ'] else 'EFFECTIVE'
    m=re.search(r'(\d{2})_(\d{2})_(\d{4})',file)
    if m: return f'{m[3]}-{m[2]}-{m[1]}', 'Datum primjene iz naziva datoteke.', 'EFFECTIVE'
    m=re.search(r'(\d{2})(\d{2})(20\d{2})',file)
    if m: return f'{m[3]}-{m[2]}-{m[1]}', 'Datum primjene iz naziva datoteke.', 'EFFECTIVE'
    return None, 'Datum primjene nije naveden; posljednja dostavljena referenca, bez potvrde aktualnosti.', 'UNCONFIRMED'

def records_for(source, sheet, rows):
    fam=source['family'];code_col=0;name_col=1;header=0;regex=None
    if fam in ['OLL','DLL']:
        if sheet=='legenda':
            return [(values[0], ' '.join(values[1:]), i+1, ['Oznaka','Opis'], [values[0],' '.join(values[1:])], {'kind':'legend'}) for i,r in enumerate(rows) if len(values:=[text(v) for v in r if text(v)])>1]
        regex=r'[A-Z]\d{2}[A-Z]{2}\d{2}\s+\d{3}' if sheet!='indikacije_smjernice' else r'[A-Za-z]{1,3}\d{1,4}[a-z]?'
        if sheet.startswith(('OLL','DLL')) or sheet=='Popis PSL':name_col=5
        elif sheet=='OL-magistralni pripravci':name_col=2
        else:name_col=1
    elif fam=='DTP-PZZ':header=next(i for i,r in enumerate(rows) if text(r[0])=='Šifra DTP');name_col=next(i for i,v in enumerate(rows[header]) if text(v)=='Naziv');regex=r'[A-Z]+\d+'
    elif fam=='DTP-SKZZ':header=1;regex=r'[A-Z]+\d+'
    elif fam=='DTS-SKUPINE':header=next(i for i,r in enumerate(rows) if len(r)>1 and text(r[1]).startswith('Šifra DTS'));code_col=1;name_col=2;regex=r'[A-Z0-9]{4}'
    elif fam=='DBL':header=next(i for i,r in enumerate(rows) if len(r)>1 and text(r[1])=='Šifra');code_col=1;name_col=2;regex=r'DBL\d+'
    elif fam=='NPP':header=next(i for i,r in enumerate(rows) if text(r[0])=='Šifra');regex=r'NPP\d+'
    elif fam=='KRV':header=1;regex=r'\d+'
    elif fam=='AKTIVNOSTI':header=2;code_col=3;name_col=1;regex=r'\d{1,7}'
    elif fam=='NASELJA':code_col=3;name_col=2;regex=r'\d{6}'
    elif fam=='OBLOGE':code_col=6;name_col=4;regex=r'\d{13}'
    elif fam=='DZO':code_col=1;name_col=0;regex=r'\d{8}'
    elif fam=='KATEGORIJE':header=1;regex=r'[A-Z]'
    elif fam=='DTS-POSTUPCI':regex=r'\d{5}-\d{2}'
    elif fam=='ZDRPOS':regex=r'\d{4,6}'
    elif fam in ['DJELATNOSTI','DJELATNOSTI-NEDATIRANO']:regex=r'\d{7}'
    elif fam=='SUDJELOVANJE':regex=r'[A-Z]\d{3}'
    elif fam=='OSLOBADANJE-MKB':regex=r'[A-Z]\d{2}(?:\.\d+[A-Z]?)?\*?'
    elif fam=='PODRUCJA':regex=r'[a-z]{3}\d+'
    elif fam=='GRESKE-RACUNA':regex=r'\d+'
    elif fam=='RAZDOZ':regex=r'[A-Z]\d+'
    elif fam=='GRESKE-PZZ':regex=r'.+'
    if fam=='DTP-PZZ' and sheet=='Legenda': return []
    if fam=='USTANOVE':
        result=[]
        for i,row in enumerate(rows[2:],start=3):
            for start in [0,9]:
                code=text(row[start+1]);name=text(row[start+2])
                if re.fullmatch(r'\d{9}',code) and name:
                    result.append((code,name,i,rows[1][start:start+9],row[start:start+9],{'level':'parent' if start==0 else 'branch','parent_code':text(row[1]) if start else None,'address':text(row[start+3])+' '+text(row[start+4]),'postal_code':text(row[start+5]),'city':text(row[start+6]),'institution_type_code':text(row[start+7]),'institution_type':text(row[start+8])}))
        return result
    result=[]
    for i,r in enumerate(rows[header+1:],start=header+2):
        if len(r)<=max(code_col,name_col):continue
        code=text(r[code_col]);name=text(r[name_col])
        if not code or (regex and not re.fullmatch(regex,code)):continue
        if (fam=='KRV' and code=='0') or (fam=='AKTIVNOSTI' and text(r[name_col])=='1'):continue # column number row
        if fam in ['OLL','DLL'] and not name:name=text(r[2]) or text(r[6]) # unbranded hospital items remain reference entries
        if not name:raise ValueError(f'Missing name: {source["file"]}:{sheet}:{i}:{code}')
        result.append((code,name,i,rows[header],r,{}))
    return result

def extract(input_dir, output):
    sources=[]
    for p in sorted(input_dir.rglob('*')):
        if p.suffix.lower() not in ['.xls','.xlsx'] or not identify(p.name):continue
        family,name=identify(p.name);sheets=load(p);date,basis,kind=date_info(p.name,family,sheets)
        sources.append({'path':str(p),'file':p.name,'family':family,'name':name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'source_date':date,'date_basis':basis,'date_kind':kind,'sheets':sheets})
    sources.sort(key=lambda s:(bool(re.search(r' \(\d\)\.xlsx$',s['file'])),s['file']))
    hashes={};unique=[]
    for s in sources:
        if s['sha256'] in hashes:s['disposition']='DUPLICATE';s['duplicate_of']=hashes[s['sha256']]['file']
        else:hashes[s['sha256']]=s;unique.append(s)
    families=collections.defaultdict(list)
    for s in unique:families[s['family']].append(s)
    chosen=[]
    for family,items in families.items():
        eligible=[s for s in items if not s['source_date'] or s['source_date']<=TODAY.isoformat()]
        selected=max(eligible,key=lambda s:s['source_date'] or '')
        for s in items:s['disposition']='SELECTED' if s is selected else 'OLDER'
        selected['status']='HISTORICAL' if family in ['ZDRPOS','DTS-POSTUPCI'] else ('UNCONFIRMED' if not selected['source_date'] else 'LATEST_SUPPLIED')
        chosen.append(selected)
    catalogs=[]
    for s in sorted(chosen,key=lambda s:s['family']):
        s['source_id']='HZZO-'+s['family']+'-'+s['sha256'][:12];s['entry_count']=0
        for sheet,rows in s['sheets']:
            if sheet=='Legenda' and s['family']=='DTP-PZZ':continue
            sub=s['family'] in ['OLL','DLL','PODRUCJA','GRESKE-PZZ']
            lid='HZZO:'+s['family']+(':'+sheet if sub else '')
            collected=collections.defaultdict(list)
            parsed=records_for(s,sheet,rows)
            used_rows={r[2] for r in parsed}
            notes='\n'.join(' | '.join(text(v) for v in row if v not in [None,'']) for index,row in enumerate(rows,start=1) if index not in used_rows and any(v not in [None,''] for v in row))
            if s['family']=='DTP-PZZ':
                notes+='\n'+'\n'.join(' | '.join(text(v) for v in row if v not in [None,'']) for title,legendrows in s['sheets'] if title=='Legenda' for row in legendrows)
            for code,name,row,headers,values,extra in parsed:
                record={'sheet':sheet,'row':row,'name':name,'fields':{text(h) or f'Kolona {i+1}':v for i,(h,v) in enumerate(zip(headers,values)) if v not in [None,'']},**extra}
                collected[code].append(record)
            if not collected:continue
            entries=[]
            for code,recs in sorted(collected.items()):
                # Retain contextual variants (e.g. multiple errors with the same code), never overwrite them.
                distinct=[];seen=set()
                for r in recs:
                    identity=json.dumps({k:v for k,v in r.items() if k not in ['row','sheet']},sort_keys=True,ensure_ascii=False)
                    if identity not in seen:distinct.append(r);seen.add(identity)
                description='\n'.join(' · '.join(f'{k}: {text(v)}' for k,v in r['fields'].items()) for r in distinct)
                entries.append({'code':code,'name':distinct[0]['name'],'description':description,'metadata':{'records':distinct}})
            catalogs.append({'id':lid,'source_id':s['source_id'],'name':s['name']+(' — '+sheet if sub else ''),'oid':None,'source_text':f"Datoteka: {s['file']}\nRadni list: {sheet}\nDatum: {s['source_date'] or 'nije potvrđen'}\n{s['date_basis']}\nStatus: {s['status']}\nSvi izvorni stupci i varijante redaka sačuvani su u podacima.\nIzvorne napomene/zaglavlja:\n{notes}",'external_only':False,'metadata':{'family':s['family'],'sheet':sheet,'source_date':s['source_date'],'date_basis':s['date_basis'],'date_kind':s['date_kind'],'status':s['status']},'entries':entries})
            s['entry_count']+=len(entries)
    output.mkdir(parents=True,exist_ok=True)
    def metadata(s):return {k:v for k,v in s.items() if k not in ['path','sheets']}
    manifest={'as_of':TODAY.isoformat(),'uploaded_files':len(sources),'duplicate_files':sum(s['disposition']=='DUPLICATE' for s in sources),'selected_files':len(chosen),'catalog_count':len(catalogs),'entry_count':sum(len(c['entries']) for c in catalogs),'files':[metadata(s) for s in sources]}
    (output/'hzzo-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    with (output/'hzzo.json.gz').open('wb') as f:
        with gzip.GzipFile(fileobj=f,mode='wb',mtime=0) as z:z.write(json.dumps(catalogs,ensure_ascii=False,separators=(',',':')).encode())
    report=['# Odabir dostavljenih HZZO šifrarnika','',f'Presjek: {TODAY:%d. %m. %Y.} Pregledano {len(sources)} datoteka; {manifest["duplicate_files"]} identičnih dodatnih kopija. Odabrano {len(chosen)} izvora, {len(catalogs)} tablica, {format(manifest["entry_count"],",").replace(",",".")} jedinstvenih parova šifrarnik/šifra.','', 'Datumi vrijede za dostavljene verzije; nije potvrđeno da HZZO nakon njih nije objavio novija izdanja. Datum revizije ne mijenja datum početka primjene. OLL i DLL su različite liste. DTS skupine nisu DTS postupci. Sve izvorne varijante redaka i dodatni stupci ostaju sačuvani.','', '| Šifrarnik | Odabrana datoteka | Datum | Status | Broj šifri |','|---|---|---|---|---|']
    status={'LATEST_SUPPLIED':'Najnoviji dostavljeni','UNCONFIRMED':'Datum/aktualnost nisu potvrđeni','HISTORICAL':'Povijesna referenca'}
    for s in sorted(chosen,key=lambda s:s['family']):report.append(f'| {s["name"]} | {s["file"]} | {s["source_date"] or "—"} | {status[s["status"]]} | {s["entry_count"]} |')
    report+=['','## Što treba provjeriti ili dopuniti','','- Noviju verziju **DTS postupaka** od poslane 18. 11. 2016.; lista **DTS skupina** ne zamjenjuje taj šifrarnik.','- Ako se koriste stare šifre **zdrpos** s bodovima, aktualnu zamjenu ili potvrdu primjene. Poslana oznaka upućuje na 30. 11. 2011.; moderna DTP lista ostaje zasebna.','- Datum/izdanje za **razdoz, kategorije osiguranja, aktivnosti zdravstvene zaštite, SKZZ područja rada i greške validacije**. Nedatirane liste su dostupne kao reference, bez tvrdnje da su službeno aktualne.','- **DZO_osiguravatelji_25042029_0.xls**: 2029. u nazivu nije potvrđen datum primjene. Datumi ugovora/servisa u tablici nisu datum cijele liste. Potrebna je oznaka izdanja ili poveznica na izvor.','- **Medicinske djelatnosti.xlsx** nema datum i sadrži 9 šifri koje nisu u verziji **djelzz od 1. 1. 2026.**. Obje su sačuvane odvojeno; devet šifri nije automatski dodano potvrđenoj verziji.','- Za naselja, obloge i sudjelovanje dostavljene su verzije iz 2022./2024.; ako postoje novija izdanja, treba ih poslati.','- Za puni šifrarnik ortopedskih i drugih pomagala nedostaju **Osnovna i Dopunska lista pomagala (OOP/DOP)**; poslane tablice obloga sadrže samo dimenzije i šifre obloga.','- Ako trebaju pune cijene lijekova, treba zasebna lista cijena originalnih pakiranja: dostavljena OLL nema taj stupac, a DLL sadrži doplatu, ne punu cijenu.','','## Starije verzije i duplikati','','| Datoteka | Odluka |','|---|---|']
    for s in sorted(sources,key=lambda s:s['file']):
        if s['disposition']!='SELECTED':report.append(f'| {s["file"]} | '+(f'Duplikat: {s["duplicate_of"]}' if s['disposition']=='DUPLICATE' else 'Starija verzija iste liste')+' |')
    (output/'hzzo-selection.md').write_text('\n'.join(report)+'\n')
    print(json.dumps({k:v for k,v in manifest.items() if k!='files'},ensure_ascii=False))
    for s in sorted(chosen,key=lambda s:s['family']):print(s['family'],s['source_date'],s['entry_count'],s['status'])

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('input',type=Path);parser.add_argument('--output',type=Path,default=Path('data/medical-catalogs'));args=parser.parse_args();extract(args.input,args.output)
