"""Export extracted references as text-safe XLSX and a reproducible source-data ZIP."""
import argparse,gzip,json,re,zipfile
from pathlib import Path
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font,PatternFill,Alignment
p=argparse.ArgumentParser();p.add_argument('--data',type=Path,default=Path('data/medical-catalogs'));p.add_argument('--output',type=Path,default=Path('/workspace/exports'));args=p.parse_args();args.output.mkdir(parents=True,exist_ok=True)
manifest=json.loads((args.data/'hzzo-manifest.json').read_text());lists=json.loads(gzip.decompress((args.data/'hzzo.json.gz').read_bytes()))
book=Workbook(write_only=True)
def append(sheet,values,header=False):
 cells=[]
 for value in values:
  cell=WriteOnlyCell(sheet,value=str(value) if value is not None else '');cell.data_type='s'
  if header:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='087F8C')
  cells.append(cell)
 sheet.append(cells)
s=book.create_sheet('Odabir izvora');append(s,['Datoteka','Šifrarnik','Datum izvora','Osnova datuma','Status','Odluka','Šifri','Identična kopija datoteke'],True)
for source in manifest['files']:append(s,[source['file'],source['family'],source['source_date'],source['date_basis'],source.get('status',''),source['disposition'],source.get('entry_count',''),source.get('duplicate_of','')])
for l in lists:
 name=re.sub(r'[\\/?*:\[\]]','-',l['id'].removeprefix('HZZO:'))[:31];s=book.create_sheet(name);s.freeze_panes='A2';s.column_dimensions['A'].width=25;s.column_dimensions['B'].width=80
 source=next(x for x in manifest['files']if x.get('source_id')==l['source_id']);serialized=[json.dumps(e['metadata'],ensure_ascii=False) for e in l['entries']];parts=max((len(v)+29999)//30000 for v in serialized)
 append(s,['Šifra (tekst)','Naziv','Izvor','Datum','Status']+[f'Izvorni podaci {i+1}' for i in range(parts)],True)
 for e,raw in zip(l['entries'],serialized):append(s,[e['code'],e['name'],source['file'],source['source_date'],source['status']]+[raw[i*30000:(i+1)*30000]for i in range(parts)])
xlsx=args.output/'HZZO-odabrani-sifrarnici-2026-10-11.xlsx';book.save(xlsx)
archive=args.output/'HZZO-izdvojeni-podaci-2026-10-11.zip'
with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED)as z:
 for name in ['hzzo-manifest.json','hzzo-selection.md','hzzo.json.gz']:z.write(args.data/name,arcname=name)
 z.write(xlsx,arcname=xlsx.name)
print(f'Exported {len(lists)} catalog sheets, {manifest["entry_count"]} codes: {xlsx.name}, {archive.name}')
